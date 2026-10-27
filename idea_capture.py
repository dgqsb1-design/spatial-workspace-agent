#!/usr/bin/env python3
"""
idea_capture: scan handwriting or sketches anywhere in a room (wall, mirror,
cupboard, whiteboard) on demand and save them to a local Word file and a .tex file.

  - Equations come out as ready-to-paste LaTeX.
  - Plain notes come out as text.
  - Diagrams and anything vague come out as a description in words, plus a draft
    TikZ figure when the sketch is simple enough.
  - Every capture also keeps the photo, so nothing is lost if a read is wrong.

Install:  pip install opencv-python anthropic numpy python-docx
          export ANTHROPIC_API_KEY=sk-ant-...
          (optional voice trigger)  pip install SpeechRecognition pyaudio

Run:      python idea_capture.py                      # one camera
          python idea_capture.py --camera 0 --camera http://192.168.1.20:8080/video
          python idea_capture.py --voice              # also trigger by saying "scan"

Trigger:  press SPACE in a preview window (or say "scan" with --voice). q quits.
Output:   notes/ideas.docx, notes/ideas.tex, notes/images/<timestamp>_cam<N>.jpg
"""

import argparse
import base64
import json
import os
import sys
import threading
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from anthropic import Anthropic
from docx import Document
from docx.shared import Inches, Pt

MODEL = os.environ.get("IDEA_CAPTURE_MODEL", "claude-sonnet-5-5")

PROMPT = r"""This is a photo of handwriting or sketches on a surface in a room (wall, mirror,
cupboard, whiteboard, paper). The author is a theoretical physicist jotting down ideas.
Ignore the surface itself, reflections of the room, glare, and people.
If the writing appears mirror-reversed because it is a reflection, read it as if corrected.

Return ONLY JSON, no code fences, in this form:
{"has_writing": true or false, "blocks": [ ... ]}

Each block, in reading order, is one of:
{"type": "math", "content": "complete LaTeX that can be pasted directly, including its environment (equation*, align*, ...)"}
{"type": "text", "content": "plain text of words and sentences, using inline $...$ for symbols"}
{"type": "diagram", "content": "description in words of the sketch: elements, labels, arrows, relations", "tikz": "a minimal tikzpicture environment if the sketch is simple enough to reproduce, otherwise an empty string"}
{"type": "unclear", "content": "what is there and why it cannot be read, with a best guess if you have one"}

Transcribe what is written. Do not invent content, and do not solve or correct the physics."""


# ---------------------------------------------------------------- camera ----
def open_camera(source: str) -> cv2.VideoCapture:
    cap = cv2.VideoCapture(int(source) if source.isdigit() else source)
    if not cap.isOpened():
        sys.exit(f"Cannot open camera source: {source}")
    cap.set(cv2.CAP_PROP_FRAME_WIDTH, 1920)
    cap.set(cv2.CAP_PROP_FRAME_HEIGHT, 1080)
    return cap


def sharpness(frame: np.ndarray) -> float:
    return float(cv2.Laplacian(cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY), cv2.CV_64F).var())


def grab_sharpest(cap: cv2.VideoCapture, n: int = 8) -> np.ndarray | None:
    """Read several frames and keep the least blurred one (also flushes stale buffers)."""
    frames = [f for ok, f in (cap.read() for _ in range(n)) if ok]
    return max(frames, key=sharpness) if frames else None


def enhance(frame: np.ndarray) -> np.ndarray:
    """Even out lighting and glare so faint marker strokes stay readable."""
    lab = cv2.cvtColor(frame, cv2.COLOR_BGR2LAB)
    l, a, b = cv2.split(lab)
    l = cv2.createCLAHE(clipLimit=2.5, tileGridSize=(8, 8)).apply(l)
    return cv2.cvtColor(cv2.merge((l, a, b)), cv2.COLOR_LAB2BGR)


# ----------------------------------------------------------- transcription ---
def transcribe(client: Anthropic, frame: np.ndarray) -> dict:
    ok, buf = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 92])
    if not ok:
        raise RuntimeError("JPEG encoding failed")
    resp = client.messages.create(
        model=MODEL,
        max_tokens=4000,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {
                    "type": "base64", "media_type": "image/jpeg",
                    "data": base64.b64encode(buf.tobytes()).decode()}},
                {"type": "text", "text": PROMPT},
            ],
        }],
    )
    raw = "".join(b.text for b in resp.content if b.type == "text")
    start, end = raw.find("{"), raw.rfind("}")
    if start < 0 or end < 0:
        raise ValueError("model did not return JSON")
    return json.loads(raw[start:end + 1])


# ------------------------------------------------------------------ output ---
def comment_lines(text: str) -> str:
    return "\n".join("% " + line for line in text.splitlines())


def append_tex(path: Path, stamp: str, cam: int, blocks: list):
    out = [f"\n% ===== Capture {stamp} (camera {cam}) ====="]
    for b in blocks:
        kind, content = b.get("type"), b.get("content", "").strip()
        if kind == "math":
            out.append(content)
        elif kind == "text":
            out.append(content + "\n")
        elif kind == "diagram":
            out.append(comment_lines("[diagram] " + content))
            tikz = b.get("tikz", "").strip()
            if tikz:
                out += [r"\iffalse % draft TikZ: delete this line and the matching \fi to enable",
                        r"\begin{figure}[h]\centering", tikz, r"\end{figure}", r"\fi"]
        else:
            out.append(comment_lines("[unclear] " + content))
    with path.open("a", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")


def code_paragraph(doc, text: str):
    run = doc.add_paragraph().add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def append_docx(path: Path, stamp: str, cam: int, blocks: list, image_path: Path):
    doc = Document(str(path)) if path.exists() else Document()
    doc.add_heading(f"Capture {stamp} (camera {cam})", level=2)
    doc.add_picture(str(image_path), width=Inches(3.5))
    for b in blocks:
        kind, content = b.get("type"), b.get("content", "").strip()
        if kind == "math":
            code_paragraph(doc, content)
        elif kind == "text":
            doc.add_paragraph(content)
        elif kind == "diagram":
            p = doc.add_paragraph()
            p.add_run("Diagram: ").bold = True
            p.add_run(content).italic = True
            tikz = b.get("tikz", "").strip()
            if tikz:
                doc.add_paragraph("Draft TikZ:").runs[0].italic = True
                code_paragraph(doc, tikz)
        else:
            p = doc.add_paragraph()
            p.add_run("Unclear: ").bold = True
            p.add_run(content).italic = True
    doc.save(str(path))


# ------------------------------------------------------------------- scan ----
def scan(client: Anthropic, caps: list, out_dir: Path):
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    (out_dir / "images").mkdir(parents=True, exist_ok=True)
    found = 0
    for i, cap in enumerate(caps):
        frame = grab_sharpest(cap)
        if frame is None:
            print(f"camera {i}: no frame")
            continue
        frame = enhance(frame)
        try:
            result = transcribe(client, frame)
        except Exception as e:
            print(f"camera {i}: transcription failed ({e})")
            continue
        blocks = result.get("blocks", [])
        if not result.get("has_writing") or not blocks:
            print(f"camera {i}: no writing found")
            continue
        image_path = out_dir / "images" / f"{stamp}_cam{i}.jpg"
        cv2.imwrite(str(image_path), frame)
        append_tex(out_dir / "ideas.tex", stamp, i, blocks)
        append_docx(out_dir / "ideas.docx", stamp, i, blocks, image_path)
        kinds = [b.get("type") for b in blocks]
        print(f"camera {i}: saved {len(blocks)} blocks {kinds}")
        found += 1
    if found:
        print(f"-> {out_dir / 'ideas.docx'}  and  {out_dir / 'ideas.tex'}")


def voice_listener(trigger: threading.Event, words=("scan", "capture")):
    import speech_recognition as sr
    rec, mic = sr.Recognizer(), sr.Microphone()
    with mic as src:
        rec.adjust_for_ambient_noise(src)
    print("Voice trigger on: say 'scan'.")
    while True:
        with mic as src:
            audio = rec.listen(src, phrase_time_limit=3)
        try:
            heard = rec.recognize_google(audio).lower()  # sends audio to Google's service
        except Exception:
            continue
        if any(w in heard for w in words):
            trigger.set()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--camera", action="append", default=None,
                    help="webcam index or stream URL; repeat for several cameras")
    ap.add_argument("--out", default="notes")
    ap.add_argument("--voice", action="store_true")
    args = ap.parse_args()

    client = Anthropic()  # reads ANTHROPIC_API_KEY
    caps = [open_camera(s) for s in (args.camera or ["0"])]
    out_dir = Path(args.out)
    trigger = threading.Event()
    if args.voice:
        threading.Thread(target=voice_listener, args=(trigger,), daemon=True).start()

    print("Ready. SPACE = scan, q = quit.")
    try:
        while True:
            for i, cap in enumerate(caps):
                ok, f = cap.read()
                if ok:
                    cv2.imshow(f"camera {i}", cv2.resize(f, None, fx=0.4, fy=0.4))
            key = cv2.waitKey(1) & 0xFF
            if key == ord("q"):
                break
            if key == ord(" "):
                trigger.set()
            if trigger.is_set():
                trigger.clear()
                print("Scanning...")
                scan(client, caps, out_dir)
    except KeyboardInterrupt:
        pass
    finally:
        for c in caps:
            c.release()
        cv2.destroyAllWindows()


if __name__ == "__main__":
    main()
