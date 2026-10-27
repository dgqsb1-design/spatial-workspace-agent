"""Write captures to a Word file and a .tex file."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def _comment(text: str) -> str:
    return "\n".join("% " + line for line in text.splitlines())


def append_tex(path: Path, stamp: str, label: str, blocks: list):
    out = [f"\n% ===== Capture {stamp} ({label}) ====="]
    for b in blocks:
        kind, content = b.get("type"), str(b.get("content", "")).strip()
        if kind == "math":
            out.append(content)
        elif kind == "text":
            out.append(content + "\n")
        elif kind == "diagram":
            out.append(_comment("[diagram] " + content))
            tikz = str(b.get("tikz", "")).strip()
            if tikz:
                out += [r"\iffalse % draft TikZ: delete this line and the matching \fi to enable",
                        r"\begin{figure}[h]\centering", tikz, r"\end{figure}", r"\fi"]
        else:
            out.append(_comment("[unclear] " + content))
    with path.open("a", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")


def _code(doc, text: str):
    run = doc.add_paragraph().add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def append_docx(path: Path, stamp: str, label: str, blocks: list, image_path: Path):
    doc = Document(str(path)) if path.exists() else Document()
    doc.add_heading(f"Capture {stamp} ({label})", level=2)
    doc.add_picture(str(image_path), width=Inches(3.5))
    for b in blocks:
        kind, content = b.get("type"), str(b.get("content", "")).strip()
        if kind == "math":
            _code(doc, content)
        elif kind == "text":
            doc.add_paragraph(content)
        elif kind == "diagram":
            p = doc.add_paragraph()
            p.add_run("Diagram: ").bold = True
            p.add_run(content).italic = True
            tikz = str(b.get("tikz", "")).strip()
            if tikz:
                doc.add_paragraph().add_run("Draft TikZ:").italic = True
                _code(doc, tikz)
        else:
            p = doc.add_paragraph()
            p.add_run("Unclear: ").bold = True
            p.add_run(content).italic = True
    doc.save(str(path))
